
from pathlib import Path
from tempfile import NamedTemporaryFile
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.collectors.models import RawLeadRecord


EXPORT_COLUMNS = [
    "STT",
    "company_name",
    "industry",
    "sub_industry",
    "province",
    "address",
    "website",
    "public_email",
    "public_phone",
    "description",
    "source_url",
    "collected_at",
]


class ExcelExporter:
    """Export processed lead records to an Excel .xlsx file."""

    def export(
        self,
        records: list[RawLeadRecord],
        output_path: str | Path,
        *,
        overwrite: bool = False,
    ) -> str:
        """
        Export records to an .xlsx file.

        Args:
            records: Processed lead records in the desired output order.
            output_path: Destination path for the Excel file.
            overwrite: Whether an existing destination may be replaced.

        Returns:
            The destination path as a string.

        Raises:
            TypeError: If records or an individual record has an invalid type.
            ValueError: If output_path does not use the .xlsx extension.
            FileExistsError: If the destination exists and overwrite is False.
            OSError: If the file cannot be written.
        """
        if not isinstance(records, list):
            raise TypeError("records must be a list of RawLeadRecord")

        for index, record in enumerate(records):
            if not isinstance(record, RawLeadRecord):
                raise TypeError(
                    f"records[{index}] must be a RawLeadRecord"
                )

        destination = Path(output_path)

        if destination.suffix.lower() != ".xlsx":
            raise ValueError("output_path must have a .xlsx extension")

        destination.parent.mkdir(parents=True, exist_ok=True)

        if destination.exists() and not overwrite:
            raise FileExistsError(
                f"Output file already exists: {destination}"
            )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Leads"

        # Write the fixed header.
        worksheet.append(EXPORT_COLUMNS)

        # Add basic formatting to the header.
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

        # Write records without modifying the original objects.
        for index, record in enumerate(records, start=1):
            row = [index]

            for field_name in EXPORT_COLUMNS[1:]:
                value = getattr(record, field_name)

                # None should be represented by a blank Excel cell.
                row.append(value if value is not None else None)

            worksheet.append(row)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        # Set readable default column widths.
        for column_index, column_name in enumerate(
            EXPORT_COLUMNS, start=1
        ):
            column_letter = get_column_letter(column_index)

            if column_name == "STT":
                worksheet.column_dimensions[column_letter].width = 8
            elif column_name in {"description", "address", "source_url"}:
                worksheet.column_dimensions[column_letter].width = 35
            else:
                worksheet.column_dimensions[column_letter].width = max(
                    16, len(column_name) + 2
                )

        # Save to a temporary file first, so a failed save does not
        # leave a partially written destination file.
        temp_path = None

        try:
            with NamedTemporaryFile(
                suffix=".xlsx",
                prefix=".excel_export_",
                dir=destination.parent,
                delete=False,
            ) as temp_file:
                temp_path = Path(temp_file.name)

            workbook.save(temp_path)
            workbook.close()

            # Recheck before publishing the final file.
            if destination.exists() and not overwrite:
                raise FileExistsError(
                    f"Output file already exists: {destination}"
                )

            os.replace(temp_path, destination)
            temp_path = None

        finally:
            workbook.close()

            if temp_path is not None:
                temp_path.unlink(missing_ok=True)

        return str(destination)
